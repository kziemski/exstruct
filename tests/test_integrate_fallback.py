from pathlib import Path

from openpyxl import Workbook

from exstruct.core import integrate


def test_extract_workbook_fallback_on_com_failure(monkeypatch, tmp_path: Path) -> None:
    # create a tiny workbook
    xlsx = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["A", "B"])
    ws.append([1, 2])
    from openpyxl.worksheet.table import Table

    ws.add_table(Table(displayName="Table1", ref="A1:B2"))
    wb.save(xlsx)
    wb.close()

    def boom(_path):
        raise RuntimeError("COM unavailable")

    monkeypatch.setattr(integrate, "_open_workbook", boom)
    result = integrate.extract_workbook(xlsx, mode="standard")
    assert result.sheets["Sheet1"].shapes == []
    assert result.sheets["Sheet1"].charts == []
    # table_candidates populated via openpyxl detection fallback
    assert result.sheets["Sheet1"].table_candidates
